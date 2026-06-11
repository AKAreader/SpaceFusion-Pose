# TaskFusion_dataset.py
# coding:utf-8
from __future__ import annotations

import os
import glob
from typing import Dict, List, Optional, Tuple

import numpy as np
import torch
from torch.utils.data.dataset import Dataset
from PIL import Image
import cv2

IMG_EXTS = (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff")

def _stem(p: str) -> str:
    return os.path.splitext(os.path.basename(p))[0]

def _list_images(folder: str) -> List[str]:
    files: List[str] = []
    for ext in IMG_EXTS:
        files += glob.glob(os.path.join(folder, f"*{ext}"))
        files += glob.glob(os.path.join(folder, f"*{ext.upper()}"))
    return sorted(files, key=lambda x: os.path.basename(x))

def _resolve_path(val: str, base_dir: str) -> str:
    val = str(val).strip().strip('"').strip("'")
    if not val:
        return ""
    if os.path.isabs(val) and os.path.exists(val):
        return val
    cand = os.path.join(base_dir, val)
    if os.path.exists(cand):
        return cand
    st = _stem(val)
    for ext in IMG_EXTS:
        c = os.path.join(base_dir, st + ext)
        if os.path.exists(c):
            return c
        c2 = os.path.join(base_dir, st + ext.upper())
        if os.path.exists(c2):
            return c2
    return ""

def _build_pairs_by_stem(vis_dir: str, ir_dir: str) -> List[Tuple[str, str, str]]:
    vis = _list_images(vis_dir)
    ir = _list_images(ir_dir)
    vis_map = {_stem(p).lower(): p for p in vis}
    ir_map = {_stem(p).lower(): p for p in ir}
    common = sorted(set(vis_map.keys()) & set(ir_map.keys()))
    return [(vis_map[k], ir_map[k], os.path.basename(vis_map[k])) for k in common]

def _read_excel_rows(xlsx_path: str) -> Tuple[List[str], List[List[object]]]:
    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise RuntimeError("读取 .xlsx 需要 openpyxl：pip install openpyxl") from e
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(list(r))
    wb.close()
    if not rows:
        return [], []
    header = [str(x).strip() if x is not None else "" for x in rows[0]]
    data = rows[1:]
    return header, data

def _pick_col(header: List[str], keywords: List[str]) -> Optional[int]:
    header_l = [h.strip().lower() for h in header]
    for kw in keywords:
        kw = kw.lower()
        for i, h in enumerate(header_l):
            if kw in h:
                return i
    return None

def _pose_schema(header: List[str]) -> Tuple[str, List[int]]:
    """
    Return (pose_type, indices)
      pose_type: 'quat' | 'euler' | 'none'
      indices: list of column indices in the order expected.
    """
    h = [x.strip().lower() for x in header]

    def find_exact(names: List[str]) -> Optional[int]:
        for n in names:
            n = n.lower()
            for i, hh in enumerate(h):
                if hh == n:
                    return i
        return None

    def find_contains(names: List[str]) -> Optional[int]:
        for n in names:
            n = n.lower()
            for i, hh in enumerate(h):
                if n in hh:
                    return i
        return None

    # quaternion candidates
    qw = find_exact(["qw", "q_w", "quat_w"]) or find_contains(["四元数w", "quatw", "q_w"])
    qx = find_exact(["qx", "q_x", "quat_x"]) or find_contains(["四元数x", "quatx", "q_x"])
    qy = find_exact(["qy", "q_y", "quat_y"]) or find_contains(["四元数y", "quaty", "q_y"])
    qz = find_exact(["qz", "q_z", "quat_z"]) or find_contains(["四元数z", "quatz", "q_z"])

    # alternative q0..q3
    q0 = find_exact(["q0"]) or find_contains(["q0"])
    q1 = find_exact(["q1"]) or find_contains(["q1"])
    q2 = find_exact(["q2"]) or find_contains(["q2"])
    q3 = find_exact(["q3"]) or find_contains(["q3"])

    if qw is not None and qx is not None and qy is not None and qz is not None:
        return "quat", [qw, qx, qy, qz]
    if q0 is not None and q1 is not None and q2 is not None and q3 is not None:
        # assume (w,x,y,z) = (q0,q1,q2,q3)
        return "quat", [q0, q1, q2, q3]

    # euler candidates
    roll = find_exact(["roll"]) or find_contains(["roll", "滚转", "phi"])
    pitch = find_exact(["pitch"]) or find_contains(["pitch", "俯仰", "theta"])
    yaw = find_exact(["yaw"]) or find_contains(["yaw", "偏航", "psi"])
    if roll is not None and pitch is not None and yaw is not None:
        return "euler", [roll, pitch, yaw]

    rx = find_exact(["rx"]) or find_contains(["rx"])
    ry = find_exact(["ry"]) or find_contains(["ry"])
    rz = find_exact(["rz"]) or find_contains(["rz"])
    if rx is not None and ry is not None and rz is not None:
        return "euler", [rx, ry, rz]

    return "none", []

class Fusion_dataset(Dataset):
    """
    无语义GT版本：返回 (vis_rgb, ir, pose_vec, pose_mask, name)
    - pose_vec: float32, shape [P]（若excel里找不到姿态列，则 P=1 且全0）
    - pose_mask: float32, shape [1]，有姿态=1，无姿态=0
    """
    def __init__(
        self,
        split: str,
        ir_path: str,
        vi_path: str,
        xlsx_path: Optional[str] = None,
        use_excel_mapping: bool = True,
        split_seed: int = 42,
        train_ratio: float = 0.9,
        resize_hw: Tuple[int, int] = (480, 640),  # (H,W)
    ):
        super().__init__()
        assert split in ["train", "val", "all", "test"]
        self.split = split
        self.ir_path = ir_path
        self.vi_path = vi_path
        self.xlsx_path = xlsx_path
        self.use_excel_mapping = use_excel_mapping
        self.split_seed = int(split_seed)
        self.train_ratio = float(train_ratio)
        self.resize_hw = resize_hw

        if not os.path.isdir(self.ir_path):
            raise FileNotFoundError(f"IR目录不存在：{self.ir_path}")
        if not os.path.isdir(self.vi_path):
            raise FileNotFoundError(f"RGB目录不存在：{self.vi_path}")

        self.pose_type = "none"
        self.pose_dim = 1  # default placeholder
        pairs: List[Tuple[str, str, str, Optional[np.ndarray]]] = []

        # 1) excel mapping (preferred)
        if self.xlsx_path and self.use_excel_mapping and os.path.exists(self.xlsx_path):
            header, data = _read_excel_rows(self.xlsx_path)
            if header and data:
                rgb_col = _pick_col(header, ["rgb", "vis", "visible"])
                ir_col = _pick_col(header, ["ir", "infra", "infrared"])
                id_col = _pick_col(header, ["id", "name", "filename", "file", "image"])

                pose_type, pose_idx = _pose_schema(header)
                self.pose_type = pose_type
                self.pose_dim = len(pose_idx) if pose_type != "none" else 1

                vis_map = {_stem(p).lower(): p for p in _list_images(self.vi_path)}
                ir_map = {_stem(p).lower(): p for p in _list_images(self.ir_path)}

                for row in data:
                    rgb_path = ""
                    ir_path = ""
                    if rgb_col is not None and rgb_col < len(row):
                        rgb_path = _resolve_path(row[rgb_col], self.vi_path)
                    if ir_col is not None and ir_col < len(row):
                        ir_path = _resolve_path(row[ir_col], self.ir_path)

                    if (not rgb_path or not ir_path) and id_col is not None and id_col < len(row):
                        sid = str(row[id_col]).strip()
                        if sid:
                            k = _stem(sid).lower()
                            if not rgb_path and k in vis_map:
                                rgb_path = vis_map[k]
                            if not ir_path and k in ir_map:
                                ir_path = ir_map[k]

                    if not (rgb_path and ir_path and os.path.exists(rgb_path) and os.path.exists(ir_path)):
                        continue

                    name = os.path.basename(rgb_path)

                    pose_vec: Optional[np.ndarray] = None
                    if self.pose_type != "none":
                        vals = []
                        ok = True
                        for ci in pose_idx:
                            if ci >= len(row):
                                ok = False
                                break
                            v = row[ci]
                            if v is None:
                                ok = False
                                break
                            try:
                                vals.append(float(v))
                            except Exception:
                                ok = False
                                break
                        if ok:
                            pose_vec = np.asarray(vals, dtype=np.float32)

                    pairs.append((rgb_path, ir_path, name, pose_vec))

        # 2) fallback: stem pairing
        if not pairs:
            base_pairs = _build_pairs_by_stem(self.vi_path, self.ir_path)
            pairs = [(v, i, n, None) for (v, i, n) in base_pairs]
            self.pose_type = "none"
            self.pose_dim = 1

        if not pairs:
            raise RuntimeError("未找到任何可配对的RGB/IR图像，请检查文件夹与命名规则。")

        # split train/val
        if split in ["train", "val"]:
            rng = np.random.default_rng(self.split_seed)
            idx = np.arange(len(pairs))
            rng.shuffle(idx)
            n_train = max(1, int(len(idx) * self.train_ratio))
            if split == "train":
                sel = idx[:n_train]
            else:
                sel = idx[n_train:] if n_train < len(idx) else idx[:1]
            self.pairs = [pairs[int(i)] for i in sel]
        else:
            self.pairs = pairs

        self.length = len(self.pairs)

        # pose stats (optional normalization for euler)
        self.pose_mean = None
        self.pose_std = None
        if self.pose_type == "euler":
            all_pose = [p for *_, p in self.pairs if p is not None]
            if len(all_pose) >= 10:
                mat = np.stack(all_pose, axis=0)
                self.pose_mean = mat.mean(axis=0).astype(np.float32)
                self.pose_std = (mat.std(axis=0) + 1e-6).astype(np.float32)

    def __len__(self):
        return self.length

    def _read_vis(self, path: str) -> torch.Tensor:
        img = Image.open(path).convert("RGB")
        if self.resize_hw is not None:
            H, W = self.resize_hw
            img = img.resize((W, H), resample=Image.BILINEAR)
        arr = np.asarray(img, dtype=np.float32) / 255.0  # (H,W,3)
        arr = np.transpose(arr, (2, 0, 1))  # (3,H,W)
        return torch.from_numpy(arr)

    def _read_ir(self, path: str) -> torch.Tensor:
        ir = cv2.imread(path, cv2.IMREAD_UNCHANGED)
        if ir is None:
            raise RuntimeError(f"IR读取失败：{path}")
        if ir.ndim == 3:
            ir = cv2.cvtColor(ir, cv2.COLOR_BGR2GRAY)
        if self.resize_hw is not None:
            H, W = self.resize_hw
            ir = cv2.resize(ir, (W, H), interpolation=cv2.INTER_LINEAR)
        ir = ir.astype(np.float32)
        if ir.max() > 255.0:
            ir = ir / 65535.0
        else:
            ir = ir / 255.0
        ir = np.expand_dims(ir, axis=0)  # (1,H,W)
        return torch.from_numpy(ir)

    def __getitem__(self, idx: int):
        vis_path, ir_path, name, pose = self.pairs[idx]
        vis = self._read_vis(vis_path)
        ir = self._read_ir(ir_path)

        if pose is None:
            pose_vec = torch.zeros((self.pose_dim,), dtype=torch.float32)
            pose_mask = torch.zeros((1,), dtype=torch.float32)
        else:
            pv = pose.astype(np.float32)
            if self.pose_type == "euler" and self.pose_mean is not None and self.pose_std is not None:
                pv = (pv - self.pose_mean) / self.pose_std
            pose_vec = torch.from_numpy(pv)
            pose_mask = torch.ones((1,), dtype=torch.float32)

        return vis, ir, pose_vec, pose_mask, name
